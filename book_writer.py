import argparse
import os
import re
import io
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Inches

class BookWriter:
    def __init__(self, theme, api_key=None):
        self.theme = theme
        self.api_key = api_key
        self.doc = Document()
        # Sanitize filename (still kept for CLI usage)
        safe_theme = re.sub(r'[^\w\s-]', '', self.theme).strip().replace(' ', '_')
        self.filename = f"{safe_theme}.docx"
        self._setup_kdp_format()
        self._setup_styles()
        self.client = None
        if self.api_key:
            try:
                from openai import OpenAI
                self.client = OpenAI(api_key=self.api_key)
            except ImportError:
                print("OpenAI library not installed. API features disabled.")

    def _setup_kdp_format(self):
        """Sets up the document format for Amazon KDP (6x9 inches)."""
        section = self.doc.sections[0]
        section.page_width = Inches(6)
        section.page_height = Inches(9)
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    def _setup_styles(self):
        """Configures basic styles for the document."""
        style = self.doc.styles['Normal']
        font = style.font
        font.name = 'Georgia'
        font.size = Pt(11)

    def generate_outline(self):
        """Generates a dynamic outline for the book based on the theme."""
        print(f"Gerando sumário dinâmico para: {self.theme}...")
        
        if self.client:
            try:
                prompt = f"Crie um sumário detalhado (apenas os nomes dos capítulos) para um livro sobre '{self.theme}'. Retorne apenas uma lista separada por quebras de linha."
                response = self.client.chat.completions.create(
                    model="gpt-4-turbo-preview",
                    messages=[{"role": "user", "content": prompt}]
                )
                outline = [line.strip("- ").strip() for line in response.choices[0].message.content.strip().split("\n") if line.strip()]
                return outline
            except Exception as e:
                print(f"Erro ao gerar sumário via API: {e}. Usando sumário padrão.")

        return [
            "Introdução",
            f"O que é {self.theme}",
            "História e Contexto",
            "Principais Conceitos",
            "Aplicações na Prática",
            "Desafios Atuais",
            "O Futuro do Setor",
            "Conclusão"
        ]

    def generate_content(self, section_title):
        """Generates content for a given section."""
        if self.client:
            try:
                prompt = f"Escreva um capítulo detalhado de livro sobre '{section_title}' para um livro cujo tema principal é '{self.theme}'. Use um tom profissional e envolvente. Escreva em Português."
                response = self.client.chat.completions.create(
                    model="gpt-4-turbo-preview",
                    messages=[{"role": "user", "content": prompt}]
                )
                return response.choices[0].message.content
            except Exception as e:
                print(f"Erro ao chamar API: {e}. Usando conteúdo de fallback.")

        content_map = {
            "Introdução": f"Bem-vindo a esta exploração sobre {self.theme}. Neste livro, vamos mergulhar nos detalhes que tornam este assunto tão fascinante e relevante hoje.",
            "Conclusão": f"Ao final desta jornada sobre {self.theme}, esperamos que você tenha uma compreensão clara e inspiradora sobre o futuro e as possibilidades que nos aguardam."
        }
        base_text = content_map.get(section_title, f"Este capítulo explora profundamente '{section_title}'.")
        
        return (
            f"{base_text}\n\n"
            f"No contexto de '{self.theme}', é impossível ignorar o impacto que '{section_title}' exerce sobre todos os envolvidos. "
            "Ao analisarmos os dados disponíveis, percebemos uma tendência clara de evolução e adaptação às novas demandas do mercado e da sociedade.\n\n"
            "Lorem ipsum dolor sit amet, consectetur adipiscing elit. Sed do eiusmod tempor incididunt ut labore et dolore magna aliqua. "
            "Duis aute irure dolor in reprehenderit in voluptate velit esse cillum dolore eu fugiat nulla pariatur."
        )

    def add_table_of_contents(self, outline):
        """Adds a manual Table of Contents."""
        self.doc.add_heading("Sumário", level=1)
        for section in outline:
            p = self.doc.add_paragraph(section)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        self.doc.add_page_break()

    def create_book(self):
        """Orchestrates the book creation process."""
        print(f"Iniciando a criação do livro: {self.theme}")
        
        self.doc.add_paragraph("\n" * 8)
        title = self.doc.add_heading(self.theme, 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.doc.add_page_break()

        outline = self.generate_outline()
        self.add_table_of_contents(outline)

        for section in outline:
            print(f"Gerando conteúdo para: {section}...")
            self.doc.add_heading(section, level=1)
            content = self.generate_content(section)
            p = self.doc.add_paragraph(content)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            self.doc.add_page_break()

    def save_to_buffer(self):
        """Saves the document to a BytesIO buffer for web usage."""
        buffer = io.BytesIO()
        self.doc.save(buffer)
        buffer.seek(0)
        return buffer

    def save_document(self):
        """Saves the document to a .docx file for CLI usage."""
        self.doc.save(self.filename)
        print(f"Livro salvo com sucesso: {self.filename}")

def main():
    parser = argparse.ArgumentParser(description="Escritor de livros para Amazon KDP.")
    parser.add_argument("--theme", type=str, required=True, help="O tema do livro.")
    parser.add_argument("--api-key", type=str, help="OpenAI API Key.")
    
    args = parser.parse_args()
    
    writer = BookWriter(args.theme, api_key=args.api_key or os.getenv("OPENAI_API_KEY"))
    writer.create_book()
    writer.save_document()

if __name__ == "__main__":
    main()
